# Copyright (c) 2026, Windify Technologies Pvt. Ltd.
# Lead Sales Report Exporter
#
# Generates an Excel "Bảng theo dõi doanh số" from Lead records.

import io
from datetime import date

import frappe
from frappe import _
from frappe.utils import add_days, flt, getdate


@frappe.whitelist()
def export_lead_sales_report(month=None, year=None):
	"""Generate and stream the sales tracking Excel report.

	Args:
		month: int 1-12, defaults to current month
		year:  int e.g. 2026, defaults to current year
	"""
	frappe.has_permission("Lead", "export", throw=True)

	try:
		import openpyxl
		from openpyxl.styles import (
			Alignment,
			Border,
			Font,
			PatternFill,
			Side,
		)
	except ImportError:
		frappe.throw(_("openpyxl is required. Run: pip install openpyxl"))

	today = date.today()
	month = int(month or today.month)
	year = int(year or today.year)
	date_label = f"{month:02d}/{year}"

	# ── Filter by month and year ──────────────────────────────────────────────
	from frappe.utils import get_first_day, get_last_day
	first_day = get_first_day(date(year, month, 1))
	last_day = get_last_day(date(year, month, 1))

	# ── Fetch Leads ───────────────────────────────────────────────────────────
	leads = frappe.get_all(
		"Lead",
		filters=[
			["creation", ">=", f"{first_day} 00:00:00"],
			["creation", "<=", f"{last_day} 23:59:59"],
		],
		fields=[
			"name",
			"education_student_full_name",
			"education_parent_phone_number",
			"education_address",
			"education_date_of_birth",
			"education_grade_class",
			"gender",
			"education_campaign",
			"education_auto_sales_order",
			"education_auto_quotation",
			"education_auto_opportunity",
		],
		order_by="creation asc",
	)

	if not leads:
		frappe.throw(_("Không có dữ liệu Lead để xuất."))

	# ── Fetch items ───────────────────────────────────────────────────────────
	items_by_lead = get_best_items_for_leads(leads)

	# ── Build Excel ───────────────────────────────────────────────────────────
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = f"DoanhSoThang{month:02d}{year}"  # / not allowed in sheet names

	COLS = [
		"TÊN HỌC SINH",
		"SĐT",
		"ĐỊA CHỈ",
		"NGÀY ĐĂNG KÝ",
		"NGÀY NHẬP HỌC",
		"TUỔI",
		"KHỐI / LỚP",
		"GIỚI TÍNH",
		"KHÓA HỌC",
		"NỘI DUNG ƯU ĐÃI",
		"HỌC PHÍ NIÊM YẾT (đ)",
		"HỌC PHÍ SAU GIẢM (đ)",
		"ĐÃ CỌC",
		"ĐÃ THANH TOÁN LẦN 2 (đ)",
		"NGUỒN",
		"QUÀ TẶNG",
		"DS TÍNH HOA HỒNG (đ)",
	]
	NUM_COLS = len(COLS)

	# ── Styles ────────────────────────────────────────────────────────────────
	HEADER_FONT = Font(name="Arial", bold=True, size=10, color="1F2933")
	CELL_FONT = Font(name="Arial", size=10, color="1F2933")
	TITLE_FILL = PatternFill("solid", fgColor="5B7FA6")
	HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
	ALT_FILL = PatternFill("solid", fgColor="F7FAFC")

	thin = Side(style="thin", color="B7B7B7")
	BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

	CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
	LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

	# ── Row 1: Title ─────────────────────────────────────────────────────────
	ws.row_dimensions[1].height = 32
	title_cell = ws.cell(row=1, column=1,
		value=f"WINDIFY FUTURE ACADEMY - DOANH SỐ THÁNG {date_label}")
	title_cell.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
	title_cell.alignment = CENTER
	title_cell.fill = TITLE_FILL
	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)

	# ── Row 2: Column headers ─────────────────────────────────────────────────
	ws.row_dimensions[2].height = 36
	for col_idx, col_name in enumerate(COLS, start=1):
		cell = ws.cell(row=2, column=col_idx, value=col_name)
		cell.font = HEADER_FONT
		cell.fill = HEADER_FILL
		cell.alignment = CENTER
		cell.border = BORDER

	# ── Data rows ─────────────────────────────────────────────────────────────
	row_num = 3
	for lead_idx, lead in enumerate(leads):
		lead_items = items_by_lead.get(lead.name)

		# Calculate age
		age = ""
		if lead.get("education_date_of_birth"):
			try:
				dob = getdate(lead.education_date_of_birth)
				today_d = date.today()
				age = today_d.year - dob.year - (
					(today_d.month, today_d.day) < (dob.month, dob.day)
				)
			except Exception:
				age = ""

		# Gender display
		gender_map = {"Male": "Nam", "Female": "Nữ", "Other": "Khác"}
		gender = gender_map.get(lead.get("gender") or "", lead.get("gender") or "")

		# Base row data (shared across item rows)
		base_data = [
			lead.get("education_student_full_name") or "",  # TÊN HỌC SINH
			lead.get("education_parent_phone_number") or "",  # SĐT
			lead.get("education_address") or "",  # ĐỊA CHỈ
			"",  # NGÀY ĐĂNG KÝ
			"",  # NGÀY NHẬP HỌC
			age,  # TUỔI
			lead.get("education_grade_class") or "",  # KHỐI / LỚP
			gender,  # GIỚI TÍNH
		]

		# Campaign / source name (may be a Link or Data field)
		campaign = lead.get("education_campaign") or ""

		if lead_items:
			for item in lead_items:
				discount = flt(item.get("discount_percentage") or 0)
				uu_dai = f"Giảm {discount:g}%" if discount > 0 else ""

				row_data = base_data + [
					item.get("item_name") or item.get("item_code") or "",  # KHÓA HỌC
					uu_dai,  # NỘI DUNG ƯU ĐÃI
					flt(item.get("price_list_rate")),  # HỌC PHÍ NIÊM YẾT
					flt(item.get("rate")),  # HỌC PHÍ SAU GIẢM
					"",  # ĐÃ CỌC
					"",  # ĐÃ THANH TOÁN LẦN 2
					campaign,  # NGUỒN
					"",  # QUÀ TẶNG
					"",  # DS TÍNH HOA HỒNG
				]
				_write_data_row(ws, row_num, row_data, lead_idx, CELL_FONT, ALT_FILL, BORDER, LEFT)
				# Áp dụng format tiền tệ cho cột HỌC PHÍ (cột K = 11, cột L = 12)
				ws.cell(row=row_num, column=11).number_format = '#,##0'
				ws.cell(row=row_num, column=12).number_format = '#,##0'
				row_num += 1
		else:
			# Lead with no items — one blank row
			row_data = base_data + ["", "", "", "", "", "", campaign, "", ""]
			_write_data_row(ws, row_num, row_data, lead_idx, CELL_FONT, ALT_FILL, BORDER, LEFT)
			row_num += 1

	# ── Column widths ─────────────────────────────────────────────────────────
	col_widths = [24, 14, 28, 15, 15, 7, 13, 10, 26, 22, 22, 22, 12, 24, 18, 14, 22]
	for col_idx, width in enumerate(col_widths, start=1):
		ws.column_dimensions[_col_letter(col_idx)].width = width

	# Freeze panes: keep row 1-2 and no column freeze
	ws.freeze_panes = "A3"

	# ── Stream the file ───────────────────────────────────────────────────────
	buf = io.BytesIO()
	wb.save(buf)
	buf.seek(0)
	file_content = buf.read()

	filename = f"lead_sales_report_{year}_{month:02d}.xlsx"
	frappe.local.response.filename = filename
	frappe.local.response.filecontent = file_content
	frappe.local.response.type = "download"


@frappe.whitelist()
def export_customer_statistics(from_date=None, to_date=None):
	"""Generate and stream the customer statistics Excel report."""
	frappe.has_permission("Lead", "export", throw=True)

	try:
		import openpyxl
		from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
		from openpyxl.worksheet.datavalidation import DataValidation
	except ImportError:
		frappe.throw(_("openpyxl is required. Run: pip install openpyxl"))

	from_date, to_date = normalize_date_range(from_date, to_date)
	leads = get_customer_statistics_rows(from_date=from_date, to_date=to_date)
	if not leads:
		frappe.throw(_("Không có dữ liệu Lead để xuất."))

	campaigns = get_utm_campaign_options()

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Data Khach Hang"

	columns = [
		"STT",
		"SALE",
		"TÊN PHỤ HUYNH",
		"SĐT PH",
		"TÊN HỌC VIÊN",
		"NGÀY SINH",
		"KHỐI / LỚP",
		"GIỚI TÍNH",
		"NGUỒN KHÁCH HÀNG",
		"CHI TIẾT NGUỒN",
		"KHÓA HỌC QUAN TÂM",
		"Tình trạng",
		"MỨC ĐỘ QUAN TÂM",
		"NGÀY TIẾP NHẬN",
	]
	group_headers = [
		("THÔNG TIN", 1, 2, "D9EAF7"),
		("PHỤ HUYNH", 3, 4, "FCE4D6"),
		("HỌC VIÊN", 5, 8, "E2F0D9"),
		("NGUỒN & KHÓA HỌC", 9, 11, "FFF2CC"),
		("TRẠNG THÁI", 12, 14, "EADCF8"),
	]
	num_cols = len(columns)

	thin = Side(style="thin", color="B7B7B7")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	center = Alignment(horizontal="center", vertical="center", wrap_text=True)
	left = Alignment(horizontal="left", vertical="center", wrap_text=True)

	title_fill = PatternFill("solid", fgColor="5B7FA6")
	title = "WINDIFY FUTURE ACADEMY - DATA KHÁCH HÀNG"
	date_range_title = get_date_range_title_suffix(from_date, to_date)
	if date_range_title:
		title = f"{title} - {date_range_title}"

	title_cell = ws.cell(row=1, column=1, value=title)
	title_cell.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
	title_cell.fill = title_fill
	title_cell.alignment = center
	title_cell.border = border
	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
	ws.row_dimensions[1].height = 34

	for label, start_col, end_col, color in group_headers:
		fill = PatternFill("solid", fgColor=color)
		ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
		cell = ws.cell(row=2, column=start_col, value=label)
		cell.font = Font(name="Arial", bold=True, size=11, color="1F2933")
		cell.fill = fill
		cell.alignment = center
		for col_idx in range(start_col, end_col + 1):
			header_cell = ws.cell(row=2, column=col_idx)
			header_cell.fill = fill
			header_cell.border = border

	group_fill_by_col = {}
	for _label, start_col, end_col, color in group_headers:
		for col_idx in range(start_col, end_col + 1):
			group_fill_by_col[col_idx] = PatternFill("solid", fgColor=color)

	ws.row_dimensions[2].height = 28
	ws.row_dimensions[3].height = 42
	for col_idx, column in enumerate(columns, start=1):
		cell = ws.cell(row=3, column=col_idx, value=column)
		cell.font = Font(name="Arial", bold=True, size=10, color="1F2933")
		cell.fill = group_fill_by_col[col_idx]
		cell.alignment = center
		cell.border = border

	add_campaign_validation(wb, ws, campaigns, start_row=4, end_row=max(len(leads) + 20, 2000), column=9)

	for row_idx, lead in enumerate(leads, start=4):
		row_values = [
			row_idx - 3,
			lead.get("sale") or "",
			lead.get("parent_name") or "",
			lead.get("parent_phone") or "",
			lead.get("student_name") or "",
			lead.get("date_of_birth") or "",
			lead.get("grade_class") or "",
			get_gender_label(lead.get("gender")),
			lead.get("campaign") or "",
			lead.get("campaign_description") or "",
			lead.get("items") or "",
			"",
			lead.get("interest_level") or "",
			lead.get("creation_date") or "",
		]
		for col_idx, value in enumerate(row_values, start=1):
			cell = ws.cell(row=row_idx, column=col_idx, value=value)
			cell.font = Font(name="Arial", size=10, color="1F2933")
			cell.alignment = center if col_idx in (1, 6, 8, 14) else left
			cell.border = border
			if row_idx % 2 == 0:
				cell.fill = PatternFill("solid", fgColor="F7FAFC")

		for date_col in (6, 14):
			ws.cell(row=row_idx, column=date_col).number_format = "dd/mm/yyyy"

	col_widths = [7, 22, 24, 16, 24, 14, 14, 12, 22, 28, 32, 18, 18, 16]
	for col_idx, width in enumerate(col_widths, start=1):
		ws.column_dimensions[_col_letter(col_idx)].width = width

	ws.freeze_panes = "A4"
	ws.auto_filter.ref = f"A3:{_col_letter(num_cols)}{len(leads) + 3}"

	buf = io.BytesIO()
	wb.save(buf)
	buf.seek(0)

	filename_suffix = get_date_range_filename_suffix(from_date, to_date)
	frappe.local.response.filename = f"thong_ke_khach_hang_{filename_suffix}.xlsx"
	frappe.local.response.filecontent = buf.read()
	frappe.local.response.type = "download"


def get_customer_statistics_rows(from_date=None, to_date=None):
	filters = []
	if from_date:
		filters.append(["creation", ">=", f"{from_date} 00:00:00"])
	if to_date:
		filters.append(["creation", "<", f"{add_days(to_date, 1)} 00:00:00"])

	leads = frappe.get_all(
		"Lead",
		filters=filters,
		fields=[
			"name",
			"creation",
			"education_lead_owner_employee",
			"lead_owner",
			"education_parent_full_name",
			"education_parent_phone_number",
			"education_student_full_name",
			"education_date_of_birth",
			"education_grade_class",
			"education_interest_level",
			"gender",
			"education_campaign",
			"education_auto_opportunity",
			"education_auto_quotation",
			"education_auto_sales_order",
		],
		order_by="creation asc",
	)
	if not leads:
		return []

	lead_names = [lead.name for lead in leads]
	employee_names = get_employee_names([lead.education_lead_owner_employee for lead in leads])
	campaign_descriptions = get_campaign_descriptions([lead.education_campaign for lead in leads])
	items_by_lead = get_items_by_lead(lead_names)

	rows = []
	for lead in leads:
		rows.append(
			{
				"lead": lead.name,
				"opportunity": lead.education_auto_opportunity,
				"quotation": lead.education_auto_quotation,
				"sales_order": lead.education_auto_sales_order,
				"sale": employee_names.get(lead.education_lead_owner_employee)
				or lead.education_lead_owner_employee
				or lead.lead_owner,
				"parent_name": lead.education_parent_full_name,
				"parent_phone": lead.education_parent_phone_number,
				"student_name": lead.education_student_full_name,
				"date_of_birth": getdate(lead.education_date_of_birth) if lead.education_date_of_birth else "",
				"grade_class": lead.education_grade_class,
				"interest_level": lead.education_interest_level,
				"gender": lead.gender,
				"campaign": lead.education_campaign,
				"campaign_description": campaign_descriptions.get(lead.education_campaign),
				"items": "\n".join(items_by_lead.get(lead.name, [])),
				"creation_date": getdate(lead.creation) if lead.creation else "",
			}
		)
	return rows


def normalize_date_range(from_date=None, to_date=None):
	from_date = getdate(from_date) if from_date else None
	to_date = getdate(to_date) if to_date else None

	if from_date and to_date and from_date > to_date:
		frappe.throw(_("Từ ngày không được lớn hơn Đến ngày."))

	return from_date, to_date


def get_date_range_filename_suffix(from_date=None, to_date=None):
	if from_date and to_date:
		return f"{from_date.isoformat()}_{to_date.isoformat()}"
	if from_date:
		return f"tu_{from_date.isoformat()}"
	if to_date:
		return f"den_{to_date.isoformat()}"
	return date.today().isoformat()


def get_date_range_title_suffix(from_date=None, to_date=None):
	if from_date and to_date:
		return f"TỪ {from_date.strftime('%d/%m/%Y')} ĐẾN {to_date.strftime('%d/%m/%Y')}"
	if from_date:
		return f"TỪ {from_date.strftime('%d/%m/%Y')}"
	if to_date:
		return f"ĐẾN {to_date.strftime('%d/%m/%Y')}"
	return ""


def get_employee_names(employees):
	employees = [employee for employee in employees if employee]
	if not employees:
		return {}

	return {
		employee.name: employee.employee_name
		for employee in frappe.get_all(
			"Employee",
			filters={"name": ["in", employees]},
			fields=["name", "employee_name"],
		)
	}


def get_campaign_descriptions(campaigns):
	campaigns = [campaign for campaign in campaigns if campaign]
	if not campaigns:
		return {}

	description_field = get_campaign_description_field()
	return {
		campaign.name: campaign.get(description_field)
		for campaign in frappe.get_all(
			"UTM Campaign",
			filters={"name": ["in", campaigns]},
			fields=["name", description_field],
		)
	}


def get_campaign_description_field():
	meta = frappe.get_meta("UTM Campaign")
	if meta.has_field("campaign_description"):
		return "campaign_description"
	if meta.has_field("description"):
		return "description"
	return "name"


def get_utm_campaign_options():
	description_field = get_campaign_description_field()
	return frappe.get_all(
		"UTM Campaign",
		fields=["name", description_field],
		order_by="name asc",
	)


def get_items_by_lead(lead_names):
	child_doctype = _get_items_child_doctype()
	if not child_doctype:
		return {}

	item_fields = ["parent", "idx", "item_code", "item_name"]
	rows = frappe.get_all(
		child_doctype,
		filters={
			"parenttype": "Lead",
			"parentfield": "education_items",
			"parent": ["in", lead_names],
		},
		fields=item_fields,
		order_by="parent asc, idx asc",
	)

	items_by_lead = {}
	for row in rows:
		item_label = row.item_name or row.item_code
		if item_label:
			items_by_lead.setdefault(row.parent, []).append(item_label)
	return items_by_lead


def get_gender_label(gender):
	gender_map = {"Male": "Nam", "Female": "Nữ", "Other": "Khác"}
	return gender_map.get(gender or "", gender or "")


def add_campaign_validation(wb, ws, campaigns, start_row, end_row, column):
	if not campaigns:
		return

	from openpyxl.worksheet.datavalidation import DataValidation

	hidden = wb.create_sheet("_campaign_options")
	hidden.sheet_state = "hidden"
	for row_idx, campaign in enumerate(campaigns, start=1):
		hidden.cell(row=row_idx, column=1, value=campaign.name)
		hidden.cell(row=row_idx, column=2, value=campaign.get(get_campaign_description_field()) or "")

	last_row = len(campaigns)
	dv = DataValidation(
		type="list",
		formula1=f"='_campaign_options'!$A$1:$A${last_row}",
		allow_blank=True,
	)
	ws.add_data_validation(dv)
	dv.add(f"{_col_letter(column)}{start_row}:{_col_letter(column)}{end_row}")


def get_best_items_for_leads(leads):
	items_by_lead = {}

	def fetch_items(doctype, names_list, lead_map):
		if not names_list: return
		rows = frappe.get_all(
			f"{doctype} Item",
			filters={"parenttype": doctype, "parent": ["in", names_list]},
			fields=["parent", "item_code", "item_name", "discount_percentage", "idx", "price_list_rate", "rate"],
			order_by="parent asc, idx asc"
		)
		for row in rows:
			lead_name = lead_map.get(row.parent)
			if lead_name:
				# Set default 0 for pricing fields if missing
				row.price_list_rate = flt(row.get("price_list_rate"))
				row.rate = flt(row.get("rate"))
				row.discount_percentage = flt(row.get("discount_percentage"))
				items_by_lead.setdefault(lead_name, []).append(row)

	# 1. Sales Order
	so_to_lead = {l.education_auto_sales_order: l.name for l in leads if l.get("education_auto_sales_order")}
	fetch_items("Sales Order", list(so_to_lead.keys()), so_to_lead)

	# 2. Quotation
	quo_to_lead = {l.education_auto_quotation: l.name for l in leads if not l.get("education_auto_sales_order") and l.get("education_auto_quotation")}
	fetch_items("Quotation", list(quo_to_lead.keys()), quo_to_lead)

	# 3. Opportunity
	opp_to_lead = {l.education_auto_opportunity: l.name for l in leads if not l.get("education_auto_sales_order") and not l.get("education_auto_quotation") and l.get("education_auto_opportunity")}
	fetch_items("Opportunity", list(opp_to_lead.keys()), opp_to_lead)

	# 4. Lead (fallback)
	lead_names = [l.name for l in leads if not l.get("education_auto_sales_order") and not l.get("education_auto_quotation") and not l.get("education_auto_opportunity")]
	child_doctype = _get_items_child_doctype()
	if child_doctype and lead_names:
		rows = frappe.get_all(
			child_doctype,
			filters={"parenttype": "Lead", "parentfield": "education_items", "parent": ["in", lead_names]},
			fields=["parent", "item_code", "item_name", "discount_percentage", "idx", "price_list_rate", "rate"],
			order_by="parent asc, idx asc"
		)
		for row in rows:
			row.price_list_rate = flt(row.get("price_list_rate"))
			row.rate = flt(row.get("rate"))
			row.discount_percentage = flt(row.get("discount_percentage"))
			items_by_lead.setdefault(row.parent, []).append(row)

	return items_by_lead


# ── Helpers ───────────────────────────────────────────────────────────────────

def _get_items_child_doctype():
	"""Return the child doctype used for education_items on Lead."""
	# Check Custom Field first
	cf = frappe.db.get_value("Custom Field", "Lead-education_items", "options")
	if cf:
		return cf
	return None


def _write_data_row(ws, row_num, row_data, lead_idx, font, alt_fill, border, alignment):
	use_alt = lead_idx % 2 == 1
	for col_idx, value in enumerate(row_data, start=1):
		cell = ws.cell(row=row_num, column=col_idx, value=value)
		cell.font = font
		cell.border = border
		cell.alignment = alignment
		if use_alt:
			cell.fill = alt_fill


def _col_letter(col_idx):
	"""Convert 1-based column index to Excel letter (A, B, …, Z, AA, …)."""
	result = ""
	while col_idx:
		col_idx, remainder = divmod(col_idx - 1, 26)
		result = chr(65 + remainder) + result
	return result
